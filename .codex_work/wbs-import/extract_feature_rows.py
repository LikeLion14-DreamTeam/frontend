from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(
    r"C:\Users\nhk10\Downloads\mcm_nfc______________________________________________________________________________________v3_2026-08-11_1.xlsx"
)
OUT_DIR = Path(r"C:\h1gyeon9\Likelion14-DreamTeam\.codex_work\wbs-import")


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    wb = load_workbook(SOURCE, data_only=True)
    ws = wb["기능"]

    rows = []
    current_requirement = ""
    current_feature = ""

    for row_index in range(25, ws.max_row + 1):
        values = [clean(ws.cell(row_index, col).value) for col in range(2, 11)]
        requirement, feature, spec, description, importance, status, user_type, device_type, note = values

        if not any(values):
            continue
        if requirement:
            current_requirement = requirement
        if feature:
            current_feature = feature

        title_source = spec or feature or requirement
        if not title_source:
            continue

        rows.append(
            {
                "row": row_index,
                "requirement": current_requirement,
                "feature": current_feature,
                "spec": spec,
                "description": description,
                "importance": importance,
                "status": status,
                "user_type": user_type,
                "device_type": device_type,
                "note": note,
            }
        )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    full_path = OUT_DIR / "feature_rows_full.json"
    compact_path = OUT_DIR / "feature_rows_compact.csv"
    full_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

    with compact_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "row",
                "requirement",
                "feature",
                "spec",
                "importance",
                "status",
                "description_preview",
            ],
        )
        writer.writeheader()
        for item in rows:
            writer.writerow(
                {
                    "row": item["row"],
                    "requirement": item["requirement"],
                    "feature": item["feature"],
                    "spec": item["spec"],
                    "importance": item["importance"],
                    "status": item["status"],
                    "description_preview": item["description"][:180],
                }
            )

    compact = [
        {
            "row": item["row"],
            "requirement": item["requirement"],
            "feature": item["feature"],
            "spec": item["spec"],
            "importance": item["importance"],
            "status": item["status"],
        }
        for item in rows
    ]
    print(json.dumps({"rows": compact, "full_path": str(full_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
