from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(
    r"C:\Users\nhk10\Downloads\mcm_nfc______________________________________________________________________________________v3_2026-08-11_1.xlsx"
)


def cell_value(value):
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    wb = load_workbook(SOURCE, data_only=False)
    summary = []
    for ws in wb.worksheets:
        rows = []
        max_col = min(ws.max_column, 30)
        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, 120),
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            values = [cell_value(value) for value in row]
            if any(values):
                rows.append(values)

        summary.append(
            {
                "sheet": ws.title,
                "dimensions": ws.calculate_dimension(),
                "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
                "rows": rows,
            }
        )

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
