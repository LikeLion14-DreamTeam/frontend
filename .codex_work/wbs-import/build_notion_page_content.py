from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(
    r"C:\Users\nhk10\Downloads\mcm_nfc______________________________________________________________________________________v3_2026-08-11_1.xlsx"
)

PAGE_IDS = {
    "1.1": "a44acc3e492083f99e58018524457ac1",
    "1.2": "3baacc3e4920816aaaeeed24c8c18c63",
    "1.3": "3baacc3e49208199baaff2255e10b314",
    "2.1": "3baacc3e492081879f23e623cd7a4d6f",
    "2.2": "3baacc3e4920814abfabf9438f5c20bb",
    "2.3": "3baacc3e49208161b8b0e492bb5ee224",
    "2.4": "3baacc3e4920812da68dcbd396d09f74",
    "3.1": "3baacc3e4920812ba525d2729dbb9fb2",
    "3.2": "3baacc3e492081d788d0c8ccac837af5",
    "3.3": "3baacc3e492081388e43d72fb7719f08",
    "4.1": "3baacc3e492081749a04e966aaf279f1",
    "4.2": "3baacc3e49208191a923fb5483c973eb",
    "5.1": "3baacc3e492081b1b518c73a90075076",
    "5.2": "3baacc3e4920815d99f4c88db410f1cd",
    "5.3": "3baacc3e492081b9a0eaf140f0bdc468",
    "5.4": "3baacc3e492081549317f21df5aa4554",
    "6.1": "3baacc3e4920819aaea2f2345ad7f9c0",
    "6.2": "3baacc3e492081b58b55eaf4529a7849",
    "6.3": "3baacc3e492081fc967ffb7ce88a3772",
    "7.1": "3baacc3e492081c5accfd80eb22112ec",
    "7.2": "3baacc3e49208135afd0e87926897de1",
    "7.3": "3baacc3e492081efb877fd4eacec8fd6",
    "7.4": "3baacc3e4920810eb9f4f312f622d770",
    "7.5": "3baacc3e492081059835d51e26e28e33",
    "7.6": "3baacc3e49208135a710f69ebdf5c284",
    "8.1": "3baacc3e492081c3972ed17d2f666155",
    "8.2": "3baacc3e49208122b0acf67503297964",
}

FEATURE_ROWS = [
    31,
    33,
    34,
    36,
    37,
    38,
    39,
    41,
    42,
    43,
    45,
    46,
    48,
    51,
    56,
    57,
    59,
    61,
    62,
    64,
    65,
    66,
    67,
    68,
    69,
    71,
    72,
]

DETAIL_ROWS = [32, 49, 50, 52, 53, 54, 55, 60]

MARKDOWN_SPECIALS = "\\*~`$[]<>{}|^"


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def escape_md(text: str) -> str:
    result = []
    for char in text:
        if char in MARKDOWN_SPECIALS:
            result.append("\\" + char)
        else:
            result.append(char)
    return "".join(result)


def format_description(text: str, indent: str = "") -> str:
    text = clean(text)
    for marker in ("[기능의 근거]", "[성공 기준]", "[개발 준비 슬롯]"):
        text = text.replace(f" {marker} ", f"\n\n{marker}\n")
    paragraphs = [escape_md(part.strip()) for part in text.split("\n")]
    return "\n".join(f"{indent}{part}" if part else "" for part in paragraphs).rstrip()


def number_from_title(title: str) -> str:
    match = re.match(r"^(\d+(?:\.\d+){1,2})\s+", title)
    if not match:
        raise ValueError(f"Cannot parse WBS number from {title!r}")
    return match.group(1)


def load_rows() -> tuple[dict[str, dict], dict[str, list[dict]]]:
    wb = load_workbook(SOURCE, data_only=True)
    ws = wb.worksheets[0]

    features = {}
    details: dict[str, list[dict]] = {}

    for row_index in FEATURE_ROWS + DETAIL_ROWS:
        feature = clean(ws.cell(row_index, 3).value)
        spec = clean(ws.cell(row_index, 4).value)
        description = clean(ws.cell(row_index, 5).value)
        title = spec or feature
        number = number_from_title(title)
        item = {
            "row": row_index,
            "number": number,
            "title": title,
            "description": description,
        }
        if row_index in FEATURE_ROWS:
            features[number] = item
        else:
            parent_number = ".".join(number.split(".")[:2])
            details.setdefault(parent_number, []).append(item)

    missing_page_ids = sorted(set(features) - set(PAGE_IDS))
    if missing_page_ids:
        raise ValueError(f"Missing page ids for: {missing_page_ids}")

    return features, details


def build_content(feature: dict, detail_items: list[dict]) -> str:
    lines = [
        "## 설명 및 상세 요구사항",
        format_description(feature["description"]),
    ]

    if detail_items:
        lines.extend(["", "## 상세 기능 체크리스트"])
        for detail in detail_items:
            detail_title = re.sub(r"^\d+(?:\.\d+){2}\s+", "", detail["title"])
            lines.append(
                f"- [ ] {detail['number']} {escape_md(detail_title)} (기능명세 row {detail['row']})"
            )
            lines.extend(
                [
                    "<details>",
                    "<summary>설명 및 상세 요구사항</summary>",
                    format_description(detail["description"], indent="\t"),
                    "</details>",
                ]
            )

    return "\n".join(lines).rstrip() + "\n"


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    features, details = load_rows()
    payloads = []

    for number in sorted(features, key=lambda value: [int(part) for part in value.split(".")]):
        feature = features[number]
        content = build_content(feature, details.get(number, []))
        payloads.append(
            {
                "number": number,
                "row": feature["row"],
                "title": feature["title"],
                "page_id": PAGE_IDS[number],
                "detail_count": len(details.get(number, [])),
                "content": content,
            }
        )

    if len(sys.argv) > 1:
        requested = set(sys.argv[1:])
        payloads = [payload for payload in payloads if payload["number"] in requested]

    print(json.dumps(payloads, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
